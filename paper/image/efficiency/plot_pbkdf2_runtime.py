"""Plot PBKDF2 execution time from the measured experiment workbook."""

from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
from openpyxl import load_workbook


OUTPUT_DIR = Path(__file__).resolve().parent
WORKBOOK = OUTPUT_DIR.parent / "experiments_result.xlsx"
HEADERS = ("内部迭代次数", "Solidity 时间 ms", "原生预编译时间 ms", "EvoCrypt 时间 ms")
SERIES = (
    ("Solidity", "#566477", "s"),
    ("Native precompile", "#C56528", "^"),
    ("EvoCrypt", "#087F8C", "o"),
)


def load_measurements():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = next((sheet for sheet in workbook if sheet.title.startswith("PBKDF2 执行时间对比")), None)
    if sheet is None or tuple(sheet.cell(1, col).value for col in range(1, 5)) != HEADERS:
        raise ValueError("PBKDF2 execution-time sheet or column headers are missing")
    rows = [tuple(sheet.cell(row, col).value for col in range(1, 5))
            for row in range(2, sheet.max_row + 1)]
    workbook.close()
    if len(rows) != 7 or any(any(not isinstance(value, (int, float)) for value in row)
                             for row in rows):
        raise ValueError("Expected seven complete numeric execution-time observations")
    iterations = [row[0] for row in rows]
    if iterations != sorted(set(iterations)) or not 0 < min(iterations):
        raise ValueError("Iteration counts must be unique, positive, and increasing")
    return iterations, [[row[col] for row in rows] for col in range(1, 4)]


def main():
    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "font.size": 8,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.linewidth": 0.75,
        "legend.frameon": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
    })
    iterations, measurements = load_measurements()
    fig, ax = plt.subplots(figsize=(4.0, 2.0))
    for (label, color, marker), values in zip(SERIES, measurements):
        ax.plot(iterations, values, label=label, color=color, marker=marker,
                markersize=4, linewidth=1.65, zorder=3)
    ax.set_xscale("log")
    ax.set_xlim(0.85, 120)
    ax.set_xticks(iterations)
    ax.set_xticklabels([str(value) for value in iterations])
    ax.minorticks_off()
    ax.set_ylim(0, 9.1)
    ax.set_xlabel("PBKDF2 iterations (log scale)")
    ax.set_ylabel("Execution time (ms)")
    ax.grid(axis="y", color="#E4E8EC", linewidth=0.7, zorder=0)
    ax.tick_params(axis="both", length=3, color="#78818B")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.04), ncol=3, fontsize=7.5,
              handlelength=1.35, handletextpad=0.35, columnspacing=0.8,
              borderaxespad=0)
    fig.subplots_adjust(left=0.145, right=0.985, bottom=0.24, top=0.75)
    fig.canvas.draw()
    stem = OUTPUT_DIR / "pbkdf2_execution_time_comparison"
    fig.savefig(f"{stem}.pdf")
    fig.savefig(f"{stem}.svg")
    fig.savefig(f"{stem}.png", dpi=300)
    fig.savefig(f"{stem}.tiff", dpi=600)
    plt.close(fig)


if __name__ == "__main__":
    main()
