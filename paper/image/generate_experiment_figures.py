#!/usr/bin/env python3
"""Generate paper figures from paper/实验数据.xlsx (four sheets)."""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_XLSX = ROOT / "实验数据.xlsx"
OUT_DIR = Path(__file__).resolve().parent

PALETTE = {
    "evocrypt": "#0F4D92",
    "solidity": "#767676",
    "native": "#42949E",
}

ALGO_DISPLAY = {
    "Add": "Add",
    "Sha256": "SHA-256",
    "Blake2bSum256": "BLAKE2b",
    "Pbkdf2Sha256": "PBKDF2",
    "Dh2048Secret": "DH-2048",
    "PedersenCommit": "Pedersen",
    "SchnorrVerify": "Schnorr",
    "PolynomialMul": "PolyMul",
}


def apply_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans", "sans-serif"],
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "font.size": 7,
            "axes.spines.right": False,
            "axes.spines.top": False,
            "axes.linewidth": 0.8,
            "legend.frameon": False,
            "xtick.major.width": 0.8,
            "ytick.major.width": 0.8,
        }
    )


def parse_ms(value) -> float:
    if value is None:
        return float("nan")
    text = str(value).strip().lower().replace(" ", "")
    match = re.match(r"^([0-9.]+)ms$", text)
    if match:
        return float(match.group(1))
    return float(text)


def parse_number(value) -> float:
    if value is None:
        return float("nan")
    text = str(value).strip().replace(",", "")
    return float(text)


def read_sheet_rows(sheet_name: str) -> list[list]:
    wb = load_workbook(DATA_XLSX, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        cells = [cell for cell in row if cell is not None and str(cell).strip() != ""]
        if cells:
            rows.append(cells)
    wb.close()
    return rows


def load_execution_latency() -> tuple[list[str], list[float], list[float], list[float]]:
    rows = read_sheet_rows("执行效率")[1:]
    algos, upgrade, contract, precompile = [], [], [], []
    for row in rows:
        if len(row) < 4:
            continue
        algos.append(str(row[0]).strip())
        upgrade.append(parse_ms(row[1]))
        contract.append(parse_ms(row[2]))
        precompile.append(parse_ms(row[3]))
    return algos, upgrade, contract, precompile


def load_execution_gas() -> tuple[list[str], list[float], list[float]]:
    rows = read_sheet_rows("执行消耗gas")[1:]
    algos, upgrade, contract = [], [], []
    for row in rows:
        if len(row) < 3:
            continue
        algos.append(str(row[0]).strip())
        upgrade.append(parse_number(row[1]))
        contract.append(parse_number(row[2]))
    return algos, upgrade, contract


def load_upgrade_latency() -> tuple[list[str], list[float], list[float]]:
    rows = read_sheet_rows("升级效率")[1:]
    algos, upgrade, contract = [], [], []
    for row in rows:
        if len(row) < 3:
            continue
        algos.append(str(row[0]).strip())
        upgrade.append(parse_ms(row[1]))
        contract.append(parse_ms(row[2]))
    return algos, upgrade, contract


def load_upgrade_gas() -> tuple[list[str], list[float], list[float]]:
    rows = read_sheet_rows("升级消耗Gas")[1:]
    algos, upgrade, contract = [], [], []
    for row in rows:
        if len(row) < 3:
            continue
        algos.append(str(row[0]).strip())
        upgrade.append(parse_number(row[1]))
        contract.append(parse_number(row[2]))
    return algos, upgrade, contract


def display_labels(algos: list[str]) -> list[str]:
    return [ALGO_DISPLAY.get(a, a) for a in algos]


def grouped_bar(
    algos: list[str],
    series: list[tuple[str, list[float], str]],
    ylabel: str,
    stem: str,
    log_y: bool = False,
) -> None:
    labels = display_labels(algos)
    n = len(labels)
    n_series = len(series)
    width = 0.78 / n_series
    x = np.arange(n)

    fig_w = max(3.4, 0.42 * n + 1.2)
    fig, ax = plt.subplots(figsize=(fig_w, 2.6))

    # 对数轴上从 0 起画的 bar 会生成延伸到天外的矩形路径（log(0) 无定义），
    # 其白色描边在 PDF 几何中穿过坐标轴下方的刻度标签。改为从轴下限起画，
    # 视觉不变但路径有限。上下限按 matplotlib 对数默认边距（log 空间 5%）确定。
    bar_bottom = 0.0
    y_limits: tuple[float, float] | None = None
    if log_y:
        all_values = [v for _, vals, _ in series for v in vals]
        log_lo, log_hi = np.log10(min(all_values)), np.log10(max(all_values))
        margin = 0.05 * (log_hi - log_lo)
        y_limits = (float(10 ** (log_lo - margin)), float(10 ** (log_hi + margin)))
        bar_bottom = y_limits[0]

    for idx, (name, values, color) in enumerate(series):
        offset = (idx - (n_series - 1) / 2) * width
        heights = [v - bar_bottom for v in values] if log_y else values
        ax.bar(
            x + offset,
            heights,
            width=width,
            bottom=bar_bottom,
            label=name,
            color=color,
            edgecolor="white",
            linewidth=0.4,
            zorder=3,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=35, ha="right", rotation_mode="anchor")
    # 分类轴不需要刻度线；保留刻度线会穿过旋转后标签的包围盒。
    ax.tick_params(axis="x", which="both", length=0)
    ax.set_ylabel(ylabel)
    if log_y:
        ax.set_yscale("log")
        if y_limits is not None:
            ax.set_ylim(*y_limits)
        # 对数轴使用纯文本刻度标签：mathtext 上下标会以约 0.7 倍字号渲染，
        # 跌破 5 pt 字号下限；数据均为正，无需非正数保护。
        ax.yaxis.set_major_formatter(
            mpl.ticker.FuncFormatter(lambda v, _: f"{v / 1e6:g}M" if v >= 1e6 else f"{v / 1e3:g}k")
        )
        ax.yaxis.set_minor_formatter(mpl.ticker.NullFormatter())
    ax.grid(axis="y", color="#E6E6E6", linewidth=0.6, zorder=0)
    ax.legend(loc="upper left", bbox_to_anchor=(0, 1.02), ncol=n_series, handlelength=1.2)

    fig.tight_layout()
    save_figure(fig, stem)
    plt.close(fig)


def save_figure(fig: plt.Figure, stem: str) -> None:
    # 导出前执行多面板对齐门禁：单面板图记录 NOT APPLICABLE，
    # 多面板几何异常时阻断导出，避免交付错位版面。
    scripts_dir = os.environ.get(
        "NATURE_FIGURE_SCRIPTS",
        str(Path.home() / ".agents" / "skills" / "nature-figure" / "scripts"),
    )
    if scripts_dir not in sys.path:
        sys.path.insert(0, scripts_dir)
    try:
        from audit_panel_alignment import require_matplotlib_panel_alignment
    except ImportError:
        require_matplotlib_panel_alignment = None
    if require_matplotlib_panel_alignment is not None:
        require_matplotlib_panel_alignment(
            fig,
            json_out=OUT_DIR / f"{stem}.alignment.json",
            tolerance_pt=1.5,
            gutter_tolerance_pt=1.5,
            strict=True,
        )

    svg_path = OUT_DIR / f"{stem}.svg"
    pdf_path = OUT_DIR / f"{stem}.pdf"
    png_path = OUT_DIR / f"{stem}.png"
    fig.savefig(svg_path, bbox_inches="tight")
    fig.savefig(pdf_path, bbox_inches="tight")
    fig.savefig(png_path, dpi=600, bbox_inches="tight")
    print(f"Wrote {svg_path.name}, {pdf_path.name}, {png_path.name}")


def main() -> None:
    # 可选参数：只重新生成指定图，避免覆盖其他仍有效的实验图。
    only = sys.argv[1] if len(sys.argv) > 1 else None

    if not DATA_XLSX.exists():
        raise FileNotFoundError(f"Missing data file: {DATA_XLSX}")

    apply_style()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    # Chart 1: upgrade gas — EvoCrypt vs Solidity
    if only in (None, "upgrade-gas"):
        algos, upgrade_gas, contract_gas = load_upgrade_gas()
        grouped_bar(
            algos,
            [
                ("EvoCrypt", upgrade_gas, PALETTE["evocrypt"]),
                ("Solidity", contract_gas, PALETTE["solidity"]),
            ],
            ylabel="Upgrade gas",
            stem="figure_upgrade_gas_comparison",
            log_y=True,
        )

    # Chart 2: upgrade latency — EvoCrypt vs Solidity
    if only in (None, "upgrade-latency"):
        algos, upgrade_lat, contract_lat = load_upgrade_latency()
        grouped_bar(
            algos,
            [
                ("EvoCrypt", upgrade_lat, PALETTE["evocrypt"]),
                ("Solidity", contract_lat, PALETTE["solidity"]),
            ],
            ylabel="Upgrade latency (ms)",
            stem="figure_upgrade_latency_comparison",
            log_y=False,
        )

    # Chart 3: execution gas after upgrade — EvoCrypt vs Solidity
    if only in (None, "execution-gas"):
        algos, exec_upgrade_gas, exec_contract_gas = load_execution_gas()
        grouped_bar(
            algos,
            [
                ("EvoCrypt", exec_upgrade_gas, PALETTE["evocrypt"]),
                ("Solidity", exec_contract_gas, PALETTE["solidity"]),
            ],
            ylabel="Execution gas",
            stem="figure_execution_gas_comparison",
            log_y=True,
        )

    # Chart 4: execution latency — EvoCrypt vs Solidity vs Native
    if only in (None, "execution-latency"):
        algos, exec_upgrade, exec_contract, exec_native = load_execution_latency()
        grouped_bar(
            algos,
            [
                ("EvoCrypt", exec_upgrade, PALETTE["evocrypt"]),
                ("Solidity", exec_contract, PALETTE["solidity"]),
                ("Native algorithm", exec_native, PALETTE["native"]),
            ],
            ylabel="Execution latency (ms)",
            stem="figure_execution_latency_comparison",
            log_y=True,
        )


if __name__ == "__main__":
    main()
