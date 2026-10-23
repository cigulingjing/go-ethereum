"""Plot the PBKDF2 cumulative-gas comparison from the experiment workbook."""

from pathlib import Path
import sys

import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, MaxNLocator
from openpyxl import load_workbook


SKILL_SCRIPTS = Path.home() / ".agents/skills/nature-figure/scripts"
sys.path.insert(0, str(SKILL_SCRIPTS))
from audit_panel_alignment import require_matplotlib_panel_alignment


OUTPUT_DIR = Path(__file__).resolve().parent
WORKBOOK = OUTPUT_DIR.parent / "experiments_result.xlsx"
# 正文以整版心插入 4.8 in 宽的图，缩放比约为 1。
# 源字号取 5 pt，页上文字约为 5 pt。
FONT_PT = 5.0
STYLES = {
    "baseline": ("#087F8C", "o"),
    "lower": ("#566477", "s"),
    "higher": ("#C56528", "^"),
}


def read_block(sheet, header_row, first_row, last_row, labels):
    if [sheet.cell(header_row, col).value for col in range(1, len(labels) + 1)] != labels:
        raise ValueError(f"Unexpected column headers at row {header_row}")
    rows = [tuple(sheet.cell(row, col).value for col in range(1, len(labels) + 1))
            for row in range(first_row, last_row + 1)]
    if len(rows) != 9 or any(any(not isinstance(value, (int, float)) for value in row)
                             for row in rows):
        raise ValueError(f"Missing or nonnumeric observations in rows {first_row}:{last_row}")
    calls = [row[0] for row in rows]
    if calls != sorted(set(calls)):
        raise ValueError("Call counts must be unique and ascending")
    return calls, [[row[col] for row in rows] for col in range(1, len(labels))]


def draw_panel(ax, block, names, roles, headroom=1.16):
    calls, series = block
    for values, name, role in zip(series, names, roles):
        color, marker = STYLES[role]
        ax.plot(calls, values, color=color, marker=marker, markersize=3.8,
                linewidth=1.55, label=name, zorder=3)
    ax.set_xlim(0.5, 16.6)
    ax.set_ylim(0, max(max(values) for values in series) * headroom)
    ax.set_xticks([1, 2, 4, 6, 8, 12, 16])
    ax.yaxis.set_major_locator(MaxNLocator(nbins=5))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value / 1e6:g}"))
    ax.grid(axis="y", color="#E4E8EC", linewidth=0.7, zorder=0)
    ax.tick_params(axis="both", length=3, color="#78818B")
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.03), fontsize=FONT_PT,
              ncol=2, columnspacing=0.65, labelspacing=0.2,
              handlelength=1.4, handletextpad=0.35,
              borderaxespad=0)
    ax.set_xlabel("Calls per transaction")
    ax.set_ylabel("Cumulative gas (million)")


def main():
    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "font.size": FONT_PT,
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.linewidth": 0.75,
        "legend.frameon": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
    })
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["PBKDF2 手续费用实验"]
    baseline = read_block(sheet, 1, 2, 10,
                          ["单笔交易调用次数", "Solidity 累计 Gas", "EvoCrypt 累计 Gas"])
    gas_zero = read_block(sheet, 17, 18, 26,
                          ["调用次数n", "g₀=1500", "g₀=3000（基线）", "g₀=6000"])
    alpha = read_block(sheet, 30, 31, 39,
                       ["调用次数n", "α=1/4", "α=1/2（基线）", "α=3/4"])
    threshold = read_block(sheet, 42, 43, 51,
                           ["调用次数n", "N₀=1", "N₀=2（基线）", "N₀=4"])
    for block in (gas_zero, alpha, threshold):
        if block[0] != baseline[0] or block[1][1] != baseline[1][1]:
            raise ValueError("Baseline call counts or EvoCrypt gas do not agree across blocks")
    workbook.close()

    fig, axes = plt.subplots(2, 2, figsize=(4.8, 3.4))
    draw_panel(axes[0, 0], baseline, ["Solidity", "EvoCrypt"],
               ["lower", "baseline"])
    draw_panel(axes[0, 1], gas_zero, ["base 1500", "base 3000", "base 6000"],
               ["lower", "baseline", "higher"])
    draw_panel(axes[1, 0], alpha, ["alpha 0.25", "alpha 0.5", "alpha 0.75"],
               ["lower", "baseline", "higher"])
    draw_panel(axes[1, 1], threshold, ["threshold 1", "threshold 2", "threshold 4"],
               ["higher", "baseline", "lower"], headroom=1.25)
    fig.subplots_adjust(left=0.13, right=0.99, bottom=0.12, top=0.80,
                        wspace=0.38, hspace=0.95)
    fig.canvas.draw()
    stem = OUTPUT_DIR / "pbkdf2_cumulative_gas"
    require_matplotlib_panel_alignment(fig, json_out=f"{stem}.alignment.json",
                                       tolerance_pt=1.5, gutter_tolerance_pt=1.5,
                                       strict=True)
    fig.savefig(f"{stem}.pdf")
    fig.savefig(f"{stem}.svg")
    fig.savefig(f"{stem}.png", dpi=300)
    fig.savefig(f"{stem}.tiff", dpi=600)
    plt.close(fig)


if __name__ == "__main__":
    main()
